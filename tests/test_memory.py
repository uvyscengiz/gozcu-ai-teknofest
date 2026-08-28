"""Epizodik hafıza araması — Qdrant'a yazma ve kosinüs araması.

Testlerin yarısı mutlu yolu değil, **iki bozulma senaryosunu** koruyor:
`Gateway.embed()` kesintide `[]` döndürüyor (Görev 03) ve Qdrant'ın kendisi
erişilemez olabilir. İkisi de bir koşuyu düşürmemeli.

Her test yerel `QdrantClient(":memory:")` ile çalışıyor — ağa çıkan tek bir
test yok.
"""

from unittest.mock import Mock

from qdrant_client import QdrantClient
from qdrant_client.models import Distance, PointStruct, VectorParams

from gozcu.core.config import QDRANT_COLLECTION, QDRANT_VECTOR_SIZE
from gozcu.memory import episodic as memory
from gozcu.memory.episodic import embed_episode, point_id, search_timeline
from gozcu.core.models import Episode, Precedent
from gozcu.core.store import Store


def _client() -> QdrantClient:
    """Süreç içi Qdrant — gerçek istemcinin aynı API'si, ağ yok."""
    return QdrantClient(":memory:")


def _vec(*head: float) -> list[float]:
    """Koleksiyonun boyutuna doldurulmuş vektör.

    Boyut `config.QDRANT_VECTOR_SIZE`'dan geliyor: koleksiyonu biz kuruyoruz ve
    yanlış boyutlu bir vektör Qdrant tarafından reddedilir.
    """
    vector = [0.0] * QDRANT_VECTOR_SIZE
    for index, value in enumerate(head):
        vector[index] = value
    return vector


def _ep(summary, risk="Orta", episode_id=None, participants=()):
    return Episode(id=episode_id, start_ts=0.0, phase="outcome",
                   summary_tr=summary, preliminary_risk=risk,
                   participants=list(participants))


def _save(client, gw, *summaries):
    """Epizotları gömer; gömülmüş epizotları geri verir."""
    saved = []
    for index, summary in enumerate(summaries, 1):
        episode = _ep(summary, episode_id=index)
        embed_episode(gw, client, episode)
        saved.append(episode)
    return saved


def _points(client):
    if not client.collection_exists(QDRANT_COLLECTION):
        return []
    return client.scroll(QDRANT_COLLECTION, limit=100, with_payload=True)[0]


# --- sıralama ---------------------------------------------------------------

def test_search_ranks_the_semantically_closest_episode_first():
    client, gw = _client(), Mock()
    # sırayla: iki epizotun gömülmesi, sonra sorgunun gömülmesi
    gw.embed.side_effect = [_vec(1.0, 0.0), _vec(0.0, 1.0), _vec(0.99, 0.14)]
    _save(client, gw, "istif aracı devrildi", "personel mola verdi")

    result = search_timeline(gw, client, "araç devrilmesi")
    assert result[0].episode.summary_tr == "istif aracı devrildi"


def test_search_returns_empty_when_nothing_is_stored():
    """Boş arşiv hiçbir şey döndürmeli — operatörün bağlam değiştirdiği anın
    ortasında istisna değil. Koleksiyon henüz yokken sorgu da gömülmüyor."""
    gw = Mock()
    assert search_timeline(gw, _client(), "herhangi bir şey") == []
    gw.embed.assert_not_called()


def test_search_returns_empty_when_the_collection_holds_no_episode():
    """Koleksiyon var ama içi boş — yine sonuç yok, yine istisna yok."""
    client, gw = _client(), Mock()
    gw.embed.side_effect = [_vec(1.0, 0.0), []]
    embed_episode(gw, client, _ep("silinecek", episode_id=1))
    client.delete(QDRANT_COLLECTION, points_selector=[1])

    assert search_timeline(gw, client, "x") == []


def test_search_honours_top_k():
    client, gw = _client(), Mock()
    gw.embed.side_effect = [_vec(1.0, 0.0), _vec(0.9, 0.1), _vec(0.8, 0.2),
                            _vec(1.0, 0.0)]
    _save(client, gw, "birinci", "ikinci", "üçüncü")

    assert len(search_timeline(gw, client, "x", top_k=2)) == 2


def test_rerank_is_never_called():
    """Organizasyon `rerank`'i ÖLÇTÜ ve zararlı buldu: R@1 0,95'ten 0,55'e
    düşüyor. `Gateway.rerank` duruyor ama hafıza araması onu çağırmıyor."""
    client, gw = _client(), Mock()
    gw.embed.side_effect = [_vec(1.0, 0.0), _vec(1.0, 0.0)]
    _save(client, gw, "tek olay")

    search_timeline(gw, client, "x")
    gw.rerank.assert_not_called()


# --- yazma tarafı: `on_close` sözleşmesi ------------------------------------

def test_embed_episode_reports_true_when_a_vector_is_stored():
    client, gw = _client(), Mock()
    gw.embed.return_value = _vec(1.0, 0.0)
    episode = _ep("kaydedildi", episode_id=7, participants=["IST-04"])

    assert embed_episode(gw, client, episode) is True
    stored = _points(client)
    assert [p.id for p in stored] == [point_id(None, 7)]
    assert stored[0].payload["id"] == 7, "epizot kimliği payload'da okunabilir kalmalı"
    # Ruling 7: sonuç ikinci bir SQLite okuması olmadan kullanılabilir olmalı.
    assert stored[0].payload["summary_tr"] == "kaydedildi"
    assert stored[0].payload["preliminary_risk"] == "Orta"
    assert stored[0].payload["start_ts"] == 0.0
    assert stored[0].payload["participants"] == ["IST-04"]
    assert stored[0].payload["id"] == 7


def test_the_collection_is_created_on_the_first_write():
    """Koleksiyonu organizasyon değil BİZ kuruyoruz — kendi boyutumuzla."""
    client, gw = _client(), Mock()
    gw.embed.return_value = _vec(1.0)
    assert client.collection_exists(QDRANT_COLLECTION) is False

    embed_episode(gw, client, _ep("ilk yazma", episode_id=1))

    params = client.get_collection(QDRANT_COLLECTION).config.params.vectors
    assert params.size == QDRANT_VECTOR_SIZE
    assert params.distance == Distance.COSINE


def test_embedding_the_same_episode_twice_replaces_the_point():
    """Nokta kimliği epizot kimliği: Görev 09'un yükleyicisi ikinci kez
    çağrıldığında arşivi çoğaltmamalı."""
    client, gw = _client(), Mock()
    gw.embed.side_effect = [_vec(1.0, 0.0), _vec(0.0, 1.0)]
    episode = _ep("aynı epizot", episode_id=3)

    assert embed_episode(gw, client, episode) is True
    assert embed_episode(gw, client, episode) is True
    assert len(_points(client)) == 1, "aynı çift tek nokta bırakmalı"


def test_degraded_embed_tier_writes_nothing():
    """Bozulmuş kademe `[]` döndürüyor. Nokta yazılırsa arama onu gerçek bir
    emsalmiş gibi geri verir; boş vektöre karşı kosinüs anlamsızdır."""
    client, gw = _client(), Mock()
    gw.embed.return_value = []
    episode = _ep("bozuk kademe", episode_id=1)

    assert embed_episode(gw, client, episode) is False
    assert _points(client) == []


def test_embed_episode_rejects_a_vector_of_the_wrong_dimension():
    """Koleksiyonun boyutu sabit; farklı boyutlu bir vektör (ör. başka bir
    gömme modeline geçilmiş) yazılırsa arama tarafı patlardı."""
    client, gw = _client(), Mock()
    gw.embed.return_value = [1.0, 0.0]
    assert embed_episode(gw, client, _ep("yanlış boyut", episode_id=1)) is False
    assert _points(client) == []


def test_embed_episode_returns_false_for_an_unsaved_episode():
    """İstisna değil `False`: geri çağrı `synthesize`'ın içinden çağrılıyor,
    oradan kaçan `ValueError` epizodu ve devir teslimi birlikte götürür."""
    client = _client()
    assert embed_episode(Mock(), client, _ep("kaydedilmemiş")) is False
    assert _points(client) == []


def test_a_fallback_episode_is_not_embedded():
    """Yedek özetli bir epizot arşive GÖMÜLMEZ: arıza metni gelecekteki bir
    koşunun emsal aramasını zehirler (`search_timeline` onu gerçek bir olay
    gibi geri verirdi). Kademe düzelip özet iyileştiğinde `False` dönüşü
    çağıran tarafın yeniden gömmesine izin veriyor.
    """
    client, gw = Mock(), Mock()
    gw.embed.return_value = _vec(1.0)
    episode = Episode(id=11, start_ts=0.0, phase="outcome",
                      summary_tr="Sentez üretilemedi; ham gözlemler kayıtlı.",
                      preliminary_risk="Orta", summary_source="fallback")

    assert embed_episode(gw, client, episode) is False
    client.upsert.assert_not_called()


def test_embed_episode_never_raises_when_the_gateway_fails():
    """`on_close`'tan kaçan bir istisna, zaten yazılmış epizodu ve devir
    teslimi birlikte götürür."""
    client, gw = _client(), Mock()
    gw.embed.side_effect = RuntimeError("gateway çöktü")

    assert embed_episode(gw, client, _ep("patlayan gateway", episode_id=1)) is False
    assert _points(client) == []


# --- Qdrant erişilemezse ----------------------------------------------------

def test_embed_episode_returns_false_when_qdrant_is_unreachable():
    """Vektör veritabanının kesintisi bir koşuyu düşürmemeli — gateway ile
    aynı felsefe."""
    gw = Mock()
    gw.embed.return_value = _vec(1.0, 0.0)
    broken = Mock()
    broken.collection_exists.side_effect = ConnectionError("Connection refused")
    broken.upsert.side_effect = ConnectionError("Connection refused")

    assert embed_episode(gw, broken, _ep("erişilemez", episode_id=1)) is False


def test_search_timeline_returns_empty_when_qdrant_is_unreachable():
    gw = Mock()
    gw.embed.return_value = _vec(1.0, 0.0)
    broken = Mock()
    broken.collection_exists.side_effect = ConnectionError("Connection refused")
    broken.query_points.side_effect = ConnectionError("Connection refused")

    assert search_timeline(gw, broken, "x") == []


# --- okuma tarafı: bozuk kademeye karşı savunma -----------------------------

def test_empty_query_vector_returns_no_results():
    """Arama anında kademe bozuksa sorgu vektörü boş gelir; boş vektörle
    Qdrant'a hiç gidilmiyor."""
    client, gw = _client(), Mock()
    gw.embed.side_effect = [_vec(1.0, 0.0), []]
    _save(client, gw, "arşivdeki olay")

    assert search_timeline(gw, client, "x") == []


# --- kendi kendine eşleşme --------------------------------------------------

def test_search_excludes_the_originating_episode():
    """Görev 11 sorguyu `episode.summary_tr` ile atıyor — tam olarak gömülen
    metin. Süzülmezse risk analistinin arşiv paneli epizodu kendi emsali
    olarak gösterir."""
    client, gw = _client(), Mock()
    gw.embed.side_effect = [_vec(1.0, 0.0), _vec(0.0, 1.0), _vec(1.0, 0.0)]
    current, other = _save(client, gw, "istif aracı devrildi",
                           "personel mola verdi")

    result = search_timeline(gw, client, current.summary_tr,
                             exclude=(current.source, current.id))
    assert [p.episode.id for p in result] == [other.id]


def test_search_keeps_every_episode_when_no_exclusion_is_given():
    client, gw = _client(), Mock()
    gw.embed.side_effect = [_vec(1.0, 0.0), _vec(0.0, 1.0), _vec(1.0, 0.0)]
    current, other = _save(client, gw, "istif aracı devrildi",
                           "personel mola verdi")

    result = search_timeline(gw, client, current.summary_tr)
    assert {p.episode.id for p in result} == {current.id, other.id}


def test_search_returns_episodes_rebuilt_from_the_payload():
    """Sonuç skoruyla birlikte bir `Precedent`; epizot payload'dan kuruluyor.

    Skor `query_points` yanıtında bugün de vardı ve atılıyordu — üç tüketicisi
    var: eşik, EMSAL kartının nicel sütunu ve kalibrasyon script'i."""
    client, gw = _client(), Mock()
    gw.embed.side_effect = [_vec(1.0, 0.0), _vec(1.0, 0.0)]
    embed_episode(gw, client, Episode(id=5, start_ts=12.5, end_ts=30.0,
                                      phase="onset", summary_tr="devrilme",
                                      participants=["IST-04"],
                                      preliminary_risk="Yüksek",
                                      state="closed"))

    found = search_timeline(gw, client, "x")[0]
    assert isinstance(found, Precedent)
    assert isinstance(found.episode, Episode)
    assert (found.episode.id, found.episode.start_ts,
            found.episode.end_ts) == (5, 12.5, 30.0)
    assert (found.episode.phase, found.episode.state) == ("onset", "closed")
    assert found.episode.preliminary_risk == "Yüksek"


# --- tutamak anahtarı -------------------------------------------------------

def test_a_store_handle_is_accepted_and_indexes_per_handle():
    """`Store` tutamağı hâlâ geçerli bir ikinci argüman — ama artık bir depo
    değil, `_client()`'ın **indeks anahtarı**.

    Gömme defteri SQLite'tan silindi (arşiv yalnız Qdrant'ta yaşıyor); geriye
    kalan iddia tutamağın kendisi. Anahtar tanımlı değilken yerel istemciler
    tutamak başına bir `WeakKeyDictionary`'de tutuluyor: aynı tutamakla
    gömülen epizot bulunur, BAŞKA bir tutamakla aranırsa bulunmaz. Bu yüzden
    `load_history(gw, store)` imzasındaki `store` "kullanılmayan parametre"
    değildir ve silinemez."""
    store, gw = Store(":memory:"), Mock()
    gw.embed.return_value = _vec(1.0, 0.0)
    episode = _ep("eski çağıran")
    episode.id = 1

    assert embed_episode(gw, store, episode) is True
    assert [p.episode.summary_tr for p in search_timeline(gw, store, "x")] == \
        ["eski çağıran"]
    # Başka bir tutamak = başka bir indeks. Ölçüldü: 0 sonuç.
    assert search_timeline(gw, Store(":memory:"), "x") == []


def test_search_timeline_drops_fallback_sourced_episodes_from_earlier_runs():
    """Yazma tarafı (`embed_episode`) yedek özetli epizotları artık gömmüyor
    (bkz. `test_a_fallback_episode_is_not_embedded`) — ama team37 koleksiyonu
    KALICI ve nokta kimliği epizot kimliği: bu kısıtlamadan ÖNCE gömülmüş
    zehirli noktalar hâlâ arşivde durabilir, aynı kimlik yeniden üretilmedikçe
    üstüne yazılacakları garanti değil. Böyle bir noktayı burada DOĞRUDAN
    yazıyoruz — tam olarak bu dal öncesi koşuların yaptığı gibi — ve okuma
    tarafının onu tek başına süzdüğünü doğruluyoruz.
    """
    client, gw = _client(), Mock()
    gw.embed.side_effect = [_vec(1.0, 0.0), _vec(1.0, 0.0)]
    real = _ep("istif aracı devrildi", episode_id=2)
    embed_episode(gw, client, real)

    poisoned = Episode(id=1, start_ts=0.0, phase="outcome",
                       summary_tr="Sentez üretilemedi; ham gözlemler kayıtlı.",
                       preliminary_risk="Orta", summary_source="fallback")
    client.upsert(QDRANT_COLLECTION, points=[
        PointStruct(id=1, vector=_vec(1.0, 0.0),
                   payload=poisoned.model_dump())])

    result = search_timeline(gw, client, "x")
    assert [p.episode.id for p in result] == [2]


def test_memory_backend_reports_local_when_no_key_is_configured(monkeypatch):
    """Anahtarsız düşüş sessiz olmamalı: konsol/KPI bunu gösterebilmeli."""
    monkeypatch.setattr(memory, "QDRANT_API_KEY", "")
    assert memory.memory_backend() == "local"


def test_memory_backend_reports_qdrant_when_a_key_is_configured(monkeypatch):
    monkeypatch.setattr(memory, "QDRANT_API_KEY", "qdr-team37-test")
    assert memory.memory_backend() == "qdrant"


# --- kararlı kimlik ---------------------------------------------------------

def test_point_id_is_stable_for_the_same_source_and_episode():
    """Aynı çift → aynı nokta; ikinci gömme çoğaltmaz, üstüne yazar."""
    from gozcu.memory import point_id
    assert point_id("9f2a", 3) == point_id("9f2a", 3)


def test_point_id_separates_episodes_that_share_a_rowid():
    """B3'ün onarımı: iki farklı video da 1 numaralı epizot üretir.

    Eski kimlik (`episode.id`) ikisini TEK noktada birleştiriyordu — ölçüldü:
    iki videodan iki epizot → 1 nokta, birincisi yok oldu.
    """
    from gozcu.memory import point_id
    assert point_id("videoA", 1) != point_id("videoB", 1)


def test_video_key_reads_content_not_the_file_name(tmp_path):
    """Gradio yüklemesi iki farklı videoyu aynı isimle getirebiliyor;
    çakışma iki alakasız olayı tek noktada birleştirir — çoğaltmadan kötü."""
    from gozcu.memory import video_key
    bir, iki = tmp_path / "video.mp4", tmp_path / "başka.mp4"
    farkli = tmp_path / "farkli.mp4"
    bir.write_bytes(b"ayni icerik")
    iki.write_bytes(b"ayni icerik")
    farkli.write_bytes(b"other_ep icerik")

    assert video_key(bir) == video_key(iki), "anahtar dosya adına bakmamalı"
    assert video_key(bir) != video_key(farkli)


def test_video_key_never_raises_on_an_unreadable_path():
    """`tests/test_run.py` var olmayan bir `"video.mp4"` yolunu 29 kez
    geçiyor; atan bir `video_key` o testlerin hepsini çökertirdi."""
    from gozcu.memory import video_key
    anahtar = video_key("olmayan-bir-dosya.mp4")
    assert isinstance(anahtar, str) and anahtar


def test_a_late_episode_from_catch_up_does_not_overwrite_an_earlier_point():
    """`DecisionLoop.catch_up` ertelenmiş pencereleri sonradan işliyor ve DAHA
    ERKEN `start_ts`'li epizotlar doğurabiliyor. Kimliğin ikinci parçası bu
    yüzden `episode.id`, `start_ts` DEĞİL — zamana dayalı bir kimlik tam o
    anda iki epizodu birbirine kaydırırdı."""
    from gozcu.memory import point_id
    client, gw = _client(), Mock()
    gw.embed.return_value = _vec(1.0)
    early = _ep("önce işlenen pencere", episode_id=1)
    early.source, early.start_ts = "9f2a", 30.0
    late_ep = _ep("telafi edilen pencere", episode_id=2)
    late_ep.source, late_ep.start_ts = "9f2a", 10.0      # DAHA ERKEN damga, SONRA doğdu

    embed_episode(gw, client, early)
    embed_episode(gw, client, late_ep)

    assert len(_points(client)) == 2, "telafi epizodu öncekini ezmemeli"
    assert {p.id for p in _points(client)} == {point_id("9f2a", 1),
                                              point_id("9f2a", 2)}


def test_the_written_point_carries_the_uuid_identity():
    """Yazma ve okuma tarafı AYNI fonksiyondan gelmek zorunda."""
    from gozcu.memory import point_id
    client, gw = _client(), Mock()
    gw.embed.return_value = _vec(1.0)
    episode = _ep("devrilme", episode_id=7)
    episode.source = "9f2a"

    assert embed_episode(gw, client, episode) is True
    stored = _points(client)
    assert [p.id for p in stored] == [point_id("9f2a", 7)]
    assert stored[0].payload["id"] == 7
    assert stored[0].payload["source"] == "9f2a"


def test_a_point_written_before_the_new_fields_still_loads():
    """`Base` `extra="forbid"` ama `_episode()` bilinmeyen anahtarları süzüyor
    ve eksikler varsayılana düşüyor — çift yönlü uyumlu."""
    client, gw = _client(), Mock()
    gw.embed.side_effect = [_vec(1.0), _vec(1.0)]
    client.create_collection(
        QDRANT_COLLECTION,
        vectors_config=VectorParams(size=QDRANT_VECTOR_SIZE,
                                    distance=Distance.COSINE))
    client.upsert(QDRANT_COLLECTION, points=[PointStruct(
        id=1, vector=_vec(1.0),
        payload={"id": 1, "start_ts": 0.0, "end_ts": 4.0, "phase": "outcome",
                 "summary_tr": "eski şema kaydı", "participants": ["IST-04"],
                 "preliminary_risk": "Orta", "state": "closed"})])

    found = search_timeline(gw, client, "eski")
    assert [p.episode.summary_tr for p in found] == ["eski şema kaydı"]
    assert found[0].episode.source is None, "eksik alan varsayılana düşmeli"


# --- skor ve hesaplanan kimlikle dışlama ------------------------------------

def test_search_carries_the_score_with_each_precedent():
    """Skor jüriye görünen tek nicel işaret (EMSAL kartı) ve eşiğin ölçüldüğü
    şey. `query_points` onu zaten döndürüyordu; taşımıyorduk."""
    client, gw = _client(), Mock()
    gw.embed.side_effect = [_vec(1.0, 0.0), _vec(0.0, 1.0), _vec(1.0, 0.0)]
    _save(client, gw, "istif aracı devrildi", "personel mola verdi")

    found = search_timeline(gw, client, "araç devrilmesi")
    assert found[0].episode.summary_tr == "istif aracı devrildi"
    assert found[0].score > found[1].score
    assert 0.0 <= found[0].score <= 1.0


def test_exclusion_drops_only_the_episode_from_the_same_source():
    """Ölçülmüş tuzak: farklı videoların epizotları da 1 numarayı taşıyor;
    tamsayı kimliğe dayalı dışlama İKİSİNİ birden elerdi."""
    client, gw = _client(), Mock()
    gw.embed.side_effect = [_vec(1.0), _vec(1.0), _vec(1.0)]
    for source in ("videoA", "videoB"):
        episode = _ep("devrilme", episode_id=1)
        episode.source = source
        embed_episode(gw, client, episode)

    found = search_timeline(gw, client, "devrilme", exclude=("videoA", 1))
    assert [p.episode.source for p in found] == ["videoB"]


def test_an_open_episode_is_excluded_from_its_own_precedents():
    """`assess_risk` AÇIK epizot üzerinde koşuyor ve sorguyu tam gömülmüş
    metinle atıyor — süzülmezse epizot kendi emsali olarak listenin başına
    oturur."""
    client, gw = _client(), Mock()
    gw.embed.side_effect = [_vec(1.0), _vec(1.0)]
    open_ep = _ep("istif aracı devriliyor", episode_id=4)
    open_ep.source, open_ep.state = "9f2a", "open"
    embed_episode(gw, client, open_ep)

    found = search_timeline(gw, client, "istif aracı devriliyor",
                            exclude=("9f2a", 4))
    assert found == []


# --- eşik, kaynak tekilleştirmesi, kilit ------------------------------------

def test_a_candidate_below_the_threshold_is_dropped():
    client, gw = _client(), Mock()
    gw.embed.side_effect = [_vec(1.0, 0.0), _vec(0.0, 1.0), _vec(1.0, 0.0)]
    _save(client, gw, "istif aracı devrildi", "kantinde kuyruk uzadı")

    unfiltered = search_timeline(gw, client, "araç devrilmesi")
    assert len(unfiltered) == 2, "eşiksiz hâlde ikisi de dönmeli"

    gw.embed.side_effect = [_vec(1.0, 0.0)]
    filtered = search_timeline(gw, client, "araç devrilmesi", threshold=0.5)
    assert [p.episode.summary_tr for p in filtered] == ["istif aracı devrildi"]


def test_an_unset_threshold_is_none_not_a_zero_floor():
    """`0.0` negatif kosinüsleri süzer — yani ölçülmemiş bir eşiktir.
    Korumasız hâl `None`.

    İddia `_threshold`'ün KENDİSİNE kuruluyor, modül sabitinin o anki
    değerine değil: kalibrasyon ölçülmüş sayıları varsayılan yapacak ve
    sabite bağlı bir test o gün sessizce kırılırdı.
    """
    from gozcu.core.config import _threshold
    assert _threshold("GOZCU_OLMAYAN_BIR_ANAHTAR") is None


def test_a_configured_threshold_parses_as_a_float(monkeypatch):
    from gozcu.core.config import _threshold
    monkeypatch.setenv("GOZCU_TEST_ESIK", "0.42")
    assert _threshold("GOZCU_TEST_ESIK") == 0.42


def test_the_same_source_appears_once_and_dedup_runs_before_the_cut():
    """B8: aynı videonun ikinci koşusu emsal listesini ikizliyordu. Dedup
    `top_k` kesilmeden ÖNCE — sonra yapılırsa ikizler gerçek emsallerin
    yerini çalar."""
    client, gw = _client(), Mock()
    gw.embed.side_effect = [_vec(1.0), _vec(1.0), _vec(0.9, 0.1), _vec(1.0)]
    for episode_id in (1, 2):
        repeat_ep = _ep("aynı klibin epizodu", episode_id=episode_id)
        repeat_ep.source = "prova"
        embed_episode(gw, client, repeat_ep)
    other_ep = _ep("başka videodaki devrilme", episode_id=1)
    other_ep.source = "gercek"
    embed_episode(gw, client, other_ep)

    found = search_timeline(gw, client, "devrilme", top_k=2)
    sources = [p.episode.source for p in found]
    assert sources.count("prova") == 1, "aynı kaynak listede bir kez"
    assert "gercek" in sources, "ikiz gerçek emsalin yerini çalmamalı"


def test_sourceless_points_are_not_collapsed_into_one_bucket():
    """`None` bir kaynak DEĞİL, kaynağın yokluğu. Bu değişiklikten önce
    yazılmış her nokta onu taşıyor; hepsini tek kovaya koymak arşivi tek
    emsale indirirdi — B8'i onarırken B4'ten beter bir şey."""
    client, gw = _client(), Mock()
    gw.embed.side_effect = [_vec(1.0), _vec(0.9, 0.1), _vec(1.0)]
    _save(client, gw, "birinci eski kayıt", "ikinci eski kayıt")  # ikisi de source=None

    found = search_timeline(gw, client, "kayıt")
    assert len(found) == 2, "kimliksiz noktalar birbirini yutmamalı"


def test_concurrent_read_and_write_lose_no_result():
    """B7'nin regresyonu. Ölçüldü: `ValueError: operands could not be
    broadcast together with shapes (32,) (31,)` — `search_timeline`'ın geniş
    `except`'i onu yutuyordu ve 400 sorgunun 6'sı sessizce `[]` dönüyordu."""
    import threading
    client, gw = _client(), Mock()
    gw.embed.return_value = _vec(1.0)
    seed_ep = _ep("ilk kayıt", episode_id=1)
    seed_ep.source = "a"
    embed_episode(gw, client, seed_ep)

    empty_results = []

    def writer(n):
        episode = _ep(f"kayıt {n}", episode_id=n)
        episode.source = "a"
        embed_episode(gw, client, episode)

    def reader():
        for _ in range(20):
            if not search_timeline(gw, client, "kayıt"):
                empty_results.append(1)

    threads = [threading.Thread(target=writer, args=(n,)) for n in range(2, 12)]
    threads.append(threading.Thread(target=reader))
    for t in threads:
        t.start()
    for t in threads:
        t.join()
    assert empty_results == [], "eş zamanlı yazma okumayı sessizce boşaltmamalı"


# --- MarkItDown belge gömme (§2) -------------------------------------------

def test_embed_document_accepts_a_file_path_instead_of_bytes():
    """§2e: imza `data: bytes` yerine `file_path: Path` alır."""
    import tempfile
    from pathlib import Path
    from gozcu.core.config import QDRANT_DOCUMENT_COLLECTION
    from gozcu.memory.library import Document

    client, gw = _client(), Mock()
    gw.embed.return_value = _vec(1.0)
    doc = Document(id="abc123", name="talimat.md", size=100,
                   uploaded_at=1756368000.0)
    with tempfile.NamedTemporaryFile(mode="w", suffix=".md",
                                     delete=False) as f:
        f.write("Yangın prosedürü: alarm → tahliye → söndürme.")
        path = Path(f.name)
    try:
        result = memory.embed_document(gw, doc, path, client=client)
        assert result is True
        points = client.scroll(QDRANT_DOCUMENT_COLLECTION, limit=10,
                               with_payload=True)[0]
        assert len(points) == 1
        assert "talimat.md" in points[0].payload["name"]
    finally:
        path.unlink(missing_ok=True)


def test_embed_document_uses_markitdown_for_binary_files():
    """§2b: PDF/DOCX gibi ikili dosyalar MarkItDown ile çözülür."""
    import tempfile
    from pathlib import Path
    from unittest.mock import patch as mock_patch
    from gozcu.core.config import QDRANT_DOCUMENT_COLLECTION
    from gozcu.memory.library import Document

    client, gw = _client(), Mock()
    gw.embed.return_value = _vec(1.0)
    doc = Document(id="pdf001", name="ekipman.pdf", size=500,
                   uploaded_at=1756368000.0)
    with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as f:
        f.write(b"%PDF-1.4 fake pdf content")
        path = Path(f.name)
    try:
        mock_result = Mock()
        mock_result.text_content = "Forklift bakım kartı: fren, lastik, hidrolik."
        with mock_patch("gozcu.memory.episodic.MarkItDown") as MockMID:
            MockMID.return_value.convert.return_value = mock_result
            result = memory.embed_document(gw, doc, path, client=client)
        assert result is True
        points = client.scroll(QDRANT_DOCUMENT_COLLECTION, limit=10,
                               with_payload=True)[0]
        assert "Forklift bakım kartı" in points[0].payload["text"]
    finally:
        path.unlink(missing_ok=True)


def test_embed_document_falls_back_to_utf8_when_markitdown_fails():
    """§2b: MarkItDown başarısız → UTF-8 decode denensin."""
    import tempfile
    from pathlib import Path
    from unittest.mock import patch as mock_patch
    from gozcu.core.config import QDRANT_DOCUMENT_COLLECTION
    from gozcu.memory.library import Document

    client, gw = _client(), Mock()
    gw.embed.return_value = _vec(1.0)
    doc = Document(id="txt001", name="notlar.txt", size=50,
                   uploaded_at=1756368000.0)
    with tempfile.NamedTemporaryFile(mode="w", suffix=".txt",
                                     delete=False) as f:
        f.write("Basit metin notu.")
        path = Path(f.name)
    try:
        with mock_patch("gozcu.memory.episodic.MarkItDown") as MockMID:
            MockMID.return_value.convert.side_effect = Exception("desteklenmiyor")
            result = memory.embed_document(gw, doc, path, client=client)
        assert result is True
        points = client.scroll(QDRANT_DOCUMENT_COLLECTION, limit=10,
                               with_payload=True)[0]
        assert "Basit metin notu" in points[0].payload["text"]
    finally:
        path.unlink(missing_ok=True)


def test_embed_document_returns_false_when_both_paths_fail():
    """§2b: MarkItDown başarısız + UTF-8 başarısız → False."""
    import tempfile
    from pathlib import Path
    from unittest.mock import patch as mock_patch
    from gozcu.memory.library import Document

    client, gw = _client(), Mock()
    doc = Document(id="bin001", name="data.bin", size=50,
                   uploaded_at=1756368000.0)
    with tempfile.NamedTemporaryFile(suffix=".bin", delete=False) as f:
        f.write(bytes(range(256)))
        path = Path(f.name)
    try:
        with mock_patch("gozcu.memory.episodic.MarkItDown") as MockMID:
            MockMID.return_value.convert.side_effect = Exception("binary")
            result = memory.embed_document(gw, doc, path, client=client)
        assert result is False
    finally:
        path.unlink(missing_ok=True)


def test_embed_document_extracts_text_from_a_real_extensionless_xlsx():
    """§2b regresyonu: `markitdown[xlsx]` eklentisi kurulu değilse bu test
    kırmızı kalır. Önceki sürümde `markitdown` çıplak paket olarak eklenmişti;
    `MarkItDown().convert()` gerçek bir XLSX'te `MissingDependencyException`
    fırlatıyordu, `_extract_text`'in geniş `except`'i bunu yutup UTF-8 geri
    dönüşe düşüyordu ve o da ikili içerikte başarısız olup `False` dönüyordu
    — belge her zaman "gömülmedi" damgası yiyordu, MarkItDown hiç mock'lanmadan.

    Dosya bilerek UZANTISIZ: `library._content_path()` içeriği hep `content`
    adıyla, uzantı olmadan saklıyor; MarkItDown formatı `magika` ile
    içerikten sezmeli, dosya adından değil.
    """
    import tempfile
    from pathlib import Path
    from openpyxl import Workbook
    from gozcu.core.config import QDRANT_DOCUMENT_COLLECTION
    from gozcu.memory.library import Document

    client, gw = _client(), Mock()
    gw.embed.return_value = _vec(1.0)
    doc = Document(id="xlsx001", name="bakim-cizelgesi.xlsx", size=1,
                   uploaded_at=1756368000.0)

    workbook = Workbook()
    workbook.active["A1"] = "Forklift bakım kartı: fren, lastik, hidrolik."
    with tempfile.TemporaryDirectory() as tmp:
        # `library._content_path()` ile aynı şekil: uzantısız `content` dosyası.
        path = Path(tmp) / "content"
        workbook.save(str(path))

        result = memory.embed_document(gw, doc, path, client=client)

    assert result is True
    points = client.scroll(QDRANT_DOCUMENT_COLLECTION, limit=10,
                           with_payload=True)[0]
    assert "Forklift bakım kartı" in points[0].payload["text"]


# --- search_documents (§3) -------------------------------------------------

def test_search_documents_returns_matching_documents():
    """§3a: anlamsal arama, skor sıralı sonuç."""
    import tempfile
    from pathlib import Path
    from gozcu.core.config import QDRANT_DOCUMENT_COLLECTION
    from gozcu.core.models import DocumentResult
    from gozcu.memory.library import Document

    client, gw = _client(), Mock()
    gw.embed.side_effect = [_vec(1.0, 0.0), _vec(0.0, 1.0), _vec(0.99, 0.1)]

    for doc_id, name, text in [("d1", "vardiya.xlsx", "Gece vardiyası personeli"),
                                ("d2", "menu.txt", "Kantinde bugün mercimek")]:
        doc = Document(id=doc_id, name=name, size=50,
                       uploaded_at=1756368000.0)
        with tempfile.NamedTemporaryFile(mode="w", suffix=".txt",
                                         delete=False) as f:
            f.write(text)
            path = Path(f.name)
        try:
            memory.embed_document(gw, doc, path, client=client)
        finally:
            path.unlink(missing_ok=True)

    from gozcu.memory.episodic import search_documents
    results = search_documents(gw, "vardiya personeli", client=client)
    assert len(results) >= 1
    assert isinstance(results[0], DocumentResult)
    assert results[0].name == "vardiya.xlsx"
    assert results[0].document_id == "d1"


def test_search_documents_returns_empty_when_collection_missing():
    """Koleksiyon yokken boş liste, istisna değil."""
    from gozcu.memory.episodic import search_documents
    gw = Mock()
    result = search_documents(gw, "herhangi", client=_client())
    assert result == []
    gw.embed.assert_not_called()


def test_search_documents_honours_threshold():
    """§3c: eşik altındaki sonuçlar süzülür."""
    import tempfile
    from pathlib import Path
    from gozcu.memory.library import Document

    client, gw = _client(), Mock()
    gw.embed.side_effect = [_vec(1.0, 0.0), _vec(0.0, 1.0), _vec(1.0, 0.0)]

    for doc_id, name, text in [("d1", "a.txt", "yangın prosedürü"),
                                ("d2", "b.txt", "kantinde kuyruk")]:
        doc = Document(id=doc_id, name=name, size=50,
                       uploaded_at=1756368000.0)
        with tempfile.NamedTemporaryFile(mode="w", suffix=".txt",
                                         delete=False) as f:
            f.write(text)
            path = Path(f.name)
        try:
            memory.embed_document(gw, doc, path, client=client)
        finally:
            path.unlink(missing_ok=True)

    from gozcu.memory.episodic import search_documents
    unfiltered = search_documents(gw, "yangın", client=client)
    assert len(unfiltered) == 2

    gw.embed.side_effect = [_vec(1.0, 0.0)]
    filtered = search_documents(gw, "yangın", threshold=0.5, client=client)
    assert all(r.score >= 0.5 for r in filtered)


# --- Qdrant cleanup on delete (§4) -----------------------------------------

def test_qdrant_vector_is_deleted_when_document_is_removed():
    """§4b: silme endpoint'i Qdrant'taki vektörü de temizler."""
    import tempfile, uuid
    from pathlib import Path
    from gozcu.core.config import QDRANT_DOCUMENT_COLLECTION
    from gozcu.memory.library import Document

    client, gw = _client(), Mock()
    gw.embed.return_value = _vec(1.0)
    doc_id = uuid.uuid4().hex
    doc = Document(id=doc_id, name="silinecek.txt", size=50,
                   uploaded_at=1756368000.0)
    with tempfile.NamedTemporaryFile(mode="w", suffix=".txt",
                                     delete=False) as f:
        f.write("Bu belge silinecek.")
        path = Path(f.name)
    try:
        memory.embed_document(gw, doc, path, client=client)
    finally:
        path.unlink(missing_ok=True)

    points_before = client.scroll(QDRANT_DOCUMENT_COLLECTION, limit=10)[0]
    assert len(points_before) == 1

    from gozcu.memory.episodic import delete_document_vector
    delete_document_vector(doc_id, client=client)

    points_after = client.scroll(QDRANT_DOCUMENT_COLLECTION, limit=10)[0]
    assert len(points_after) == 0


def test_qdrant_cleanup_is_graceful_when_collection_missing():
    """§4b: koleksiyon yoksa hata değil."""
    from gozcu.memory.episodic import delete_document_vector
    delete_document_vector("nonexistent", client=_client())
